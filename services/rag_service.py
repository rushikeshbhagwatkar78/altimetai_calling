from __future__ import annotations

import json
import logging
import os
import re
import time
from typing import Any, Dict, List, Optional, Tuple

from pypdf import PdfReader
import docx

from db.models import KnowledgeChunk, KnowledgeDocument, Property
from db.session import get_db_session
from services.embedding_service import compute_cosine_similarity, generate_embedding

logger = logging.getLogger("crm_rag")

# In-memory search cache per call: (call_id, query_key) -> cached_results
CALL_SEARCH_CACHE: Dict[str, Dict[str, Any]] = {}


# ---------------------------------------------------------------------------
# Document Ingestion & Chunking
# ---------------------------------------------------------------------------
def extract_text_from_file(file_path: str, filename: str) -> str:
    """Extract raw text from uploaded document based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            pages_text = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(pages_text)
        elif ext in (".docx", ".doc"):
            doc = docx.Document(file_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paras)
        elif ext in (".txt", ".md", ".csv", ".json"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
    except Exception as e:
        logger.error(f"TEXT_EXTRACTION_ERROR | file={filename} | error={e}")
        raise

    return text.strip()


def chunk_text(text: str, chunk_size: int = 500, chunk_overlap: int = 100) -> List[str]:
    """Split text into semantic paragraphs and overlapping chunks."""
    paragraphs = re.split(r"\n{2,}", text)
    chunks = []
    current_chunk = []
    current_length = 0

    for para in paragraphs:
        cleaned_para = " ".join(para.split())
        if not cleaned_para:
            continue
        para_len = len(cleaned_para)

        if current_length + para_len > chunk_size and current_chunk:
            chunks.append(" ".join(current_chunk))
            # Keep last part for overlap
            overlap_words = " ".join(current_chunk).split()[-int(chunk_overlap / 5):]
            current_chunk = [" ".join(overlap_words), cleaned_para]
            current_length = sum(len(c) for c in current_chunk)
        else:
            current_chunk.append(cleaned_para)
            current_length += para_len

    if current_chunk:
        chunks.append(" ".join(current_chunk))

    return [c.strip() for c in chunks if len(c.strip()) > 30]


def ingest_document_file(file_path: str, filename: str, doc_id: Optional[str] = None) -> KnowledgeDocument:
    """Ingest a document, chunk it, generate embeddings, and persist to database."""
    ext = os.path.splitext(filename)[1].lower().replace(".", "")
    raw_text = extract_text_from_file(file_path, filename)
    chunks = chunk_text(raw_text)

    with get_db_session() as db:
        if doc_id:
            doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == doc_id).first()
            if not doc:
                doc = KnowledgeDocument(id=doc_id, filename=filename, document_type=ext)
                db.add(doc)
        else:
            doc = KnowledgeDocument(filename=filename, document_type=ext)
            db.add(doc)
        
        doc.status = "indexing"
        doc.chunk_count = len(chunks)
        db.flush()

        # Delete existing chunks for this doc if re-indexing
        db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == doc.id).delete()

        vector_count = 0
        for i, chunk_text_content in enumerate(chunks):
            embedding_vec = generate_embedding(chunk_text_content)
            chunk = KnowledgeChunk(
                document_id=doc.id,
                chunk_index=i,
                content=chunk_text_content,
                metadata_json=json.dumps({"filename": filename, "chunk_index": i}),
                embedding_json=json.dumps(embedding_vec),
            )
            db.add(chunk)
            vector_count += 1

        doc.vector_count = vector_count
        doc.status = "indexed"
        doc.error_message = None
        db.commit()
        db.refresh(doc)
        logger.info(f"DOCUMENT_INGESTED | id={doc.id} | file={filename} | chunks={len(chunks)}")
        return doc


# ---------------------------------------------------------------------------
# Hybrid Property Search & Knowledge Retrieval
# ---------------------------------------------------------------------------
def search_knowledge_chunks(query: str, top_k: int = 3) -> List[Dict[str, Any]]:
    """Search knowledge base chunks via vector similarity."""
    if not query.strip():
        return []

    query_vec = generate_embedding(query)
    results = []

    with get_db_session() as db:
        chunks = db.query(KnowledgeChunk).all()
        for chk in chunks:
            if not chk.embedding_json:
                continue
            try:
                chk_vec = json.loads(chk.embedding_json)
                sim = compute_cosine_similarity(query_vec, chk_vec)
                results.append({
                    "chunk_id": chk.id,
                    "document_id": chk.document_id,
                    "content": chk.content,
                    "score": round(sim, 3),
                    "metadata": json.loads(chk.metadata_json or "{}"),
                })
            except Exception:
                continue

    results.sort(key=lambda x: x["score"], reverse=True)
    return results[:top_k]


def search_properties_hybrid(
    location: Optional[str] = None,
    property_type: Optional[str] = None,
    bhk: Optional[str] = None,
    budget_max: Optional[float] = None,
    query_text: Optional[str] = None,
    call_id: Optional[str] = None,
    top_k: int = 3,
) -> Dict[str, Any]:
    """
    Hybrid Real Estate Property Search Engine:
    1. Structured SQL filter on Location, Property Type, BHK, and Budget constraints.
    2. Vector semantic scoring on preferences, amenities, and query text.
    3. Per-call in-memory caching to eliminate redundant queries.
    """
    cache_key = f"loc={location}|type={property_type}|bhk={bhk}|budget={budget_max}|q={query_text}"
    if call_id and call_id in CALL_SEARCH_CACHE:
        if cache_key in CALL_SEARCH_CACHE[call_id]:
            logger.info(f"RAG_CACHE_HIT | call_id={call_id} | key={cache_key}")
            return CALL_SEARCH_CACHE[call_id][cache_key]

    t0 = time.perf_counter()

    with get_db_session() as db:
        q = db.query(Property).filter(Property.is_active == True)

        # 1. Structured SQL filtering
        if location and location.strip():
            loc_clean = location.strip().lower()
            q = q.filter(Property.location.ilike(f"%{loc_clean}%"))

        if property_type and property_type.strip() and property_type.lower() != "any":
            type_clean = property_type.strip().lower()
            q = q.filter(Property.property_type.ilike(f"%{type_clean}%"))

        if bhk and bhk.strip():
            bhk_clean = bhk.strip().lower().replace("bhk", "").strip()
            if bhk_clean in ("1", "2", "3", "4"):
                q = q.filter(Property.bhk.ilike(f"%{bhk_clean}%"))

        if budget_max and budget_max > 0:
            # Allow properties where minimum price is within 15% of max budget
            q = q.filter(Property.price_min <= budget_max * 1.15)

        candidates = q.all()

        # If strict search yielded 0 results, relax constraints gracefully
        fallback_used = False
        if not candidates:
            fallback_used = True
            q_fallback = db.query(Property).filter(Property.is_active == True)
            if location:
                q_fallback = q_fallback.filter(Property.location.ilike(f"%{location.strip()}%"))
            candidates = q_fallback.limit(5).all()
            if not candidates:
                candidates = db.query(Property).filter(Property.is_active == True).limit(5).all()

        # 2. Ranking and scoring
        query_vector = None
        if query_text:
            query_vector = generate_embedding(query_text)

        ranked = []
        for prop in candidates:
            score = 0.85
            match_reasons = []

            if location and location.lower() in prop.location.lower():
                score += 0.05
                match_reasons.append(f"Location ({prop.location})")

            if bhk and prop.bhk and bhk.lower() in prop.bhk.lower():
                score += 0.04
                match_reasons.append(f"Configuration ({prop.bhk})")

            if budget_max and prop.price_min <= budget_max:
                score += 0.04
                match_reasons.append(f"Within Budget ({prop.price_display})")

            if query_vector and prop.description:
                prop_desc_vec = generate_embedding(f"{prop.project_name} {prop.location} {prop.amenities} {prop.description}")
                sim = compute_cosine_similarity(query_vector, prop_desc_vec)
                score = (score * 0.7) + (sim * 0.3)
                if sim > 0.6:
                    match_reasons.append("Semantic Preference Match")

            ranked.append({
                "property_id": prop.id,
                "project_name": prop.project_name,
                "location": prop.location,
                "property_type": prop.property_type,
                "bhk": prop.bhk or "N/A",
                "price_display": prop.price_display,
                "price_min": prop.price_min,
                "price_max": prop.price_max,
                "carpet_area": prop.carpet_area or "N/A",
                "possession_status": prop.possession_status,
                "builder_name": prop.builder_name,
                "description": prop.description,
                "amenities": prop.amenities or "",
                "available_units": prop.available_units,
                "match_score": round(min(score, 0.99), 2),
                "match_reasons": match_reasons or ["General Nagpur Portfolio Match"],
            })

        ranked.sort(key=lambda x: x["match_score"], reverse=True)
        top_matches = ranked[:top_k]

    elapsed_ms = round((time.perf_counter() - t0) * 1000)
    result = {
        "count": len(top_matches),
        "matches": top_matches,
        "search_latency_ms": elapsed_ms,
        "fallback_used": fallback_used,
    }

    if call_id:
        if call_id not in CALL_SEARCH_CACHE:
            CALL_SEARCH_CACHE[call_id] = {}
        CALL_SEARCH_CACHE[call_id][cache_key] = result

    logger.info(
        f"PROPERTY_SEARCH_COMPLETED | call_id={call_id} | matches={len(top_matches)} | latency={elapsed_ms}ms"
    )
    return result
